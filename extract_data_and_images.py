import openpyxl
import re
import os
import json
import time

def extract_all():
    print("Loading workbook...")
    start_time = time.time()
    excel_path = "Copy of Layout Standardization  - August 25, 12_23 PM.xlsx"
    
    wb_data = openpyxl.load_workbook(excel_path, data_only=True)
    print(f"Workbook loaded in {time.time() - start_time:.2f}s")
    
    os.makedirs("static/layout_images", exist_ok=True)
    
    # A. Map Defined Names
    defined_to_cell = {}
    for name, defn in wb_data.defined_names.items():
        ref = defn.attr_text if hasattr(defn, 'attr_text') else str(defn)
        if "Layout image" in ref:
            m = re.search(r'\$([A-Z]+)\$(\d+)', ref)
            if m:
                col_letter, row_str = m.groups()
                row = int(row_str)
                col = 0
                for ch in col_letter:
                    col = col * 26 + (ord(ch) - ord('A') + 1)
                norm_name = name.strip()
                defined_to_cell[norm_name.lower()] = (row, col)
                defined_to_cell[norm_name] = (row, col)
    
    # B. Extract images from 'Layout image' sheet
    ws_img = wb_data["Layout image"]
    img_filenames = {}
    
    for idx, img in enumerate(ws_img._images):
        from_cell = getattr(img.anchor, '_from', None)
        if from_cell:
            r = from_cell.row + 1
            c = from_cell.col + 1
            ext = img.format.lower() if img.format else "png"
            if ext == "jpeg": ext = "jpg"
            filename = f"img_r{r}_c{c}.{ext}"
            filepath = os.path.join("static/layout_images", filename)
            if not os.path.exists(filepath):
                with open(filepath, "wb") as f:
                    f.write(img._data())
            img_filenames[(r, c)] = f"/static/layout_images/{filename}"
            
    # Map row in Layout image sheet to part name in Col 1
    row_to_part = {}
    for r in range(1, 1500):
        val = ws_img.cell(r, 1).value
        if val:
            row_to_part[r] = str(val).strip()
            
    # C. Extract Data sheet rows with complete PRINT fields
    ws_data = wb_data["Data"]
    
    def clean_val(v):
        if v is None: return None
        s = str(v).strip()
        if s in ["#VALUE!", "#N/A", "#REF!", "#DIV/0!", "None", "**"]:
            return None
        return v
    
    def to_float(v, round_to=4):
        if v is None: return None
        try:
            val = float(v)
            return round(val, round_to)
        except (ValueError, TypeError):
            return None

    def to_str(v):
        if v is None: return ""
        s = str(v).strip()
        if s in ["#VALUE!", "#N/A", "#REF!", "#DIV/0!", "None", "**"]:
            return ""
        return s

    records = []
    
    for r in range(3, ws_data.max_row + 1):
        base_part = to_str(ws_data.cell(r, 1).value)
        status = to_str(ws_data.cell(r, 6).value)
        part_no = to_str(ws_data.cell(r, 7).value)
        layout_name = to_str(ws_data.cell(r, 8).value)
        rm_erp = to_str(ws_data.cell(r, 9).value)
        cut_blank = to_str(ws_data.cell(r, 10).value)
        grade = to_str(ws_data.cell(r, 11).value)
        
        if not part_no and not base_part:
            continue
            
        # Blank 1
        blank_l = to_float(ws_data.cell(r, 12).value)
        blank_w = to_float(ws_data.cell(r, 13).value)
        blank_t = to_float(ws_data.cell(r, 14).value)
        
        # RM size
        rm_l = to_float(ws_data.cell(r, 15).value)
        rm_w = to_float(ws_data.cell(r, 16).value)
        rm_t = to_float(ws_data.cell(r, 17).value)
        
        rm_weight = to_float(ws_data.cell(r, 18).value)
        blank_weight = to_float(ws_data.cell(r, 19).value)
        blank_qty = to_float(ws_data.cell(r, 20).value)
        usage_weight = to_float(ws_data.cell(r, 21).value)
        endbit_weight = to_float(ws_data.cell(r, 22).value)
        yield_val = to_float(ws_data.cell(r, 23).value)
        if yield_val is not None and yield_val <= 1.0:
            yield_pct = round(yield_val * 100, 2)
        elif yield_val is not None:
            yield_pct = round(yield_val, 2)
        else:
            yield_pct = None
            
        po_price = to_float(ws_data.cell(r, 24).value)
        sheets_used = to_float(ws_data.cell(r, 25).value) or 1.0
        wastage_cost = to_float(ws_data.cell(r, 26).value)
        
        # End bits 1 to 4
        eb1_l = to_float(ws_data.cell(r, 29).value)
        eb1_w = to_float(ws_data.cell(r, 30).value)
        eb1_q = to_float(ws_data.cell(r, 31).value)
        
        eb2_l = to_float(ws_data.cell(r, 32).value)
        eb2_w = to_float(ws_data.cell(r, 33).value)
        eb2_q = to_float(ws_data.cell(r, 34).value)
        
        eb3_l = to_float(ws_data.cell(r, 35).value)
        eb3_w = to_float(ws_data.cell(r, 36).value)
        eb3_q = to_float(ws_data.cell(r, 37).value)
        
        eb4_l = to_float(ws_data.cell(r, 38).value)
        eb4_w = to_float(ws_data.cell(r, 39).value)
        eb4_q = to_float(ws_data.cell(r, 40).value)
        
        cutting_plan = to_str(ws_data.cell(r, 41).value)
        
        # Child Parts 2, 3, 4
        part2 = to_str(ws_data.cell(r, 42).value)
        p2_l = to_float(ws_data.cell(r, 43).value)
        p2_w = to_float(ws_data.cell(r, 44).value)
        p2_t = to_float(ws_data.cell(r, 45).value)
        p2_qty = to_float(ws_data.cell(r, 46).value)
        p2_wt = to_float(ws_data.cell(r, 47).value)
        
        part3 = to_str(ws_data.cell(r, 48).value)
        p3_l = to_float(ws_data.cell(r, 49).value)
        p3_w = to_float(ws_data.cell(r, 50).value)
        p3_t = to_float(ws_data.cell(r, 51).value)
        p3_qty = to_float(ws_data.cell(r, 52).value)
        p3_wt = to_float(ws_data.cell(r, 53).value)
        
        part4 = to_str(ws_data.cell(r, 54).value)
        p4_l = to_float(ws_data.cell(r, 55).value)
        p4_w = to_float(ws_data.cell(r, 56).value)
        p4_t = to_float(ws_data.cell(r, 57).value)
        p4_qty = to_float(ws_data.cell(r, 58).value)
        p4_wt = to_float(ws_data.cell(r, 59).value)
        
        # Strip quantities
        strip_q1 = to_float(ws_data.cell(r, 60).value)
        strip_q2 = to_float(ws_data.cell(r, 61).value)
        strip_q3 = to_float(ws_data.cell(r, 62).value)
        strip_q4 = to_float(ws_data.cell(r, 63).value)
        
        # Format Cut Blank Sizes for parts 2, 3, 4
        cut_blank_2 = f"{p2_l}*{p2_w}*{p2_t}" if (p2_l and p2_w and p2_t) else ""
        cut_blank_3 = f"{p3_l}*{p3_w}*{p3_t}" if (p3_l and p3_w and p3_t) else ""
        cut_blank_4 = f"{p4_l}*{p4_w}*{p4_t}" if (p4_l and p4_w and p4_t) else ""
        
        # Find layout image
        image_url = None
        st_lower = status.lower() if status else ""
        is_non_standard = ("non standard" in st_lower or "non-standard" in st_lower or "not use" in st_lower)

        if is_non_standard:
            # For non-standard layouts: ONLY use Column 6 image if it exists for this row/part.
            # Never use standardized image (_c3) or defined names.
            # If no column 6 image exists, image_url must remain None so the drawing space stays blank.
            if part_no:
                for row_num, pname in row_to_part.items():
                    if pname.lower() == part_no.lower():
                        if (row_num, 6) in img_filenames:
                            image_url = img_filenames[(row_num, 6)]
                        break
        elif "2nd" in st_lower:
            # 2nd choice layout (Column 5)
            if layout_name:
                ln_clean = layout_name.replace(" ", "").strip()
                cell = defined_to_cell.get(ln_clean) or defined_to_cell.get(ln_clean.lower())
                if cell and cell[1] == 5 and cell in img_filenames:
                    image_url = img_filenames[cell]
            if not image_url and part_no:
                for row_num, pname in row_to_part.items():
                    if pname.lower() == part_no.lower():
                        if (row_num, 5) in img_filenames:
                            image_url = img_filenames[(row_num, 5)]
                        break
        elif "costing" in st_lower:
            # Costing layout (Column 2)
            if part_no:
                for row_num, pname in row_to_part.items():
                    if pname.lower() == part_no.lower():
                        if (row_num, 2) in img_filenames:
                            image_url = img_filenames[(row_num, 2)]
                        break
        else:
            # Standardized layout (Column 3)
            if layout_name:
                ln_clean = layout_name.replace(" ", "").strip()
                cell = defined_to_cell.get(ln_clean) or defined_to_cell.get(ln_clean.lower())
                if cell and cell in img_filenames:
                    image_url = img_filenames[cell]
            if not image_url and part_no:
                for row_num, pname in row_to_part.items():
                    if pname.lower() == part_no.lower():
                        if (row_num, 3) in img_filenames:
                            image_url = img_filenames[(row_num, 3)]
                        break
        
        rec = {
            "row_id": r,
            "base_part": base_part,
            "part_no": part_no or base_part,
            "status": status or "Unspecified",
            "layout_name": layout_name,
            "rm_erp": rm_erp,
            "cut_blank": cut_blank,
            "grade": grade,
            
            # Print Sheet Header & Parts
            "no_of_sheets": sheets_used,
            "cutting_plan": cutting_plan,
            
            # Dimensions
            "thickness": blank_t or rm_t,
            "length": rm_l,
            "width": rm_w,
            "rm_weight": rm_weight,
            
            # Part 1 Specs
            "part1": {
                "part_no": part_no,
                "cut_blank": cut_blank,
                "blank_size": {"length": blank_l, "width": blank_w, "thickness": blank_t},
                "blank_weight": blank_weight,
                "blank_qty": blank_qty,
                "strip_qty": strip_q1,
                "total_blank_qty_sheet": (blank_qty * strip_q1) if (blank_qty and strip_q1) else blank_qty,
                "total_blank_qty": ((blank_qty * strip_q1) if (blank_qty and strip_q1) else (blank_qty or 0)) * sheets_used
            },
            
            # Part 2 Specs
            "part2": {
                "part_no": part2,
                "cut_blank": cut_blank_2,
                "blank_size": {"length": p2_l, "width": p2_w, "thickness": p2_t},
                "blank_weight": p2_wt,
                "blank_qty": p2_qty,
                "strip_qty": strip_q2,
                "total_blank_qty_sheet": (p2_qty * strip_q2) if (p2_qty and strip_q2) else p2_qty,
                "total_blank_qty": ((p2_qty * strip_q2) if (p2_qty and strip_q2) else (p2_qty or 0)) * sheets_used
            } if part2 else None,

            # Part 3 Specs
            "part3": {
                "part_no": part3,
                "cut_blank": cut_blank_3,
                "blank_size": {"length": p3_l, "width": p3_w, "thickness": p3_t},
                "blank_weight": p3_wt,
                "blank_qty": p3_qty,
                "strip_qty": strip_q3,
                "total_blank_qty_sheet": (p3_qty * strip_q3) if (p3_qty and strip_q3) else p3_qty,
                "total_blank_qty": ((p3_qty * strip_q3) if (p3_qty and strip_q3) else (p3_qty or 0)) * sheets_used
            } if part3 else None,

            # Part 4 Specs
            "part4": {
                "part_no": part4,
                "cut_blank": cut_blank_4,
                "blank_size": {"length": p4_l, "width": p4_w, "thickness": p4_t},
                "blank_weight": p4_wt,
                "blank_qty": p4_qty,
                "strip_qty": strip_q4,
                "total_blank_qty_sheet": (p4_qty * strip_q4) if (p4_qty and strip_q4) else p4_qty,
                "total_blank_qty": ((p4_qty * strip_q4) if (p4_qty and strip_q4) else (p4_qty or 0)) * sheets_used
            } if part4 else None,

            # Weights & Yields
            "per_sheet": {
                "usage_weight": usage_weight,
                "endbit_weight": endbit_weight,
                "yield_pct": yield_pct
            },
            "total_sheet": {
                "usage_weight": round(usage_weight * sheets_used, 4) if usage_weight else None,
                "endbit_weight": round(endbit_weight * sheets_used, 4) if endbit_weight else None,
                "yield_pct": yield_pct
            },
            
            # End bits 1 to 4
            "endbits": [
                {"name": "End bit - 1", "dim": f"{eb1_l}*{eb1_w}*{blank_t or rm_t}" if (eb1_l and eb1_w) else "", "qty": f"{int(eb1_q)} Nos" if eb1_q else ""},
                {"name": "End bit - 2", "dim": f"{eb2_l}*{eb2_w}*{blank_t or rm_t}" if (eb2_l and eb2_w) else "", "qty": f"{int(eb2_q)} Nos" if eb2_q else ""},
                {"name": "End bit - 3", "dim": f"{eb3_l}*{eb3_w}*{blank_t or rm_t}" if (eb3_l and eb3_w) else "", "qty": f"{int(eb3_q)} Nos" if eb3_q else ""},
                {"name": "End bit - 4", "dim": f"{eb4_l}*{eb4_w}*{blank_t or rm_t}" if (eb4_l and eb4_w) else "", "qty": f"{int(eb4_q)} Nos" if eb4_q else ""}
            ],
            
            # Cost impact
            "po_price": po_price,
            "wastage_cost": wastage_cost,
            "image_url": image_url
        }
        records.append(rec)
        
    print(f"Extracted {len(records)} detailed data records")
    
    # Build Part Index
    parts_dict = {}
    base_parts_dict = {}
    
    for rec in records:
        p_no = rec["part_no"]
        b_part = rec["base_part"]
        
        if p_no not in parts_dict:
            parts_dict[p_no] = {
                "part_no": p_no,
                "base_part": b_part,
                "records": []
            }
        parts_dict[p_no]["records"].append(rec)
        
        if b_part:
            if b_part not in base_parts_dict:
                base_parts_dict[b_part] = {
                    "base_part": b_part,
                    "item_parts": set()
                }
            base_parts_dict[b_part]["item_parts"].add(p_no)
            
    for k in base_parts_dict:
        base_parts_dict[k]["item_parts"] = sorted(list(base_parts_dict[k]["item_parts"]))
        
    samples = []
    for p_no, p_data in parts_dict.items():
        has_img = any(r["image_url"] for r in p_data["records"])
        has_rm = any(r["rm_erp"] for r in p_data["records"])
        if has_img and has_rm and len(p_data["records"]) >= 2:
            samples.append({
                "part_no": p_no,
                "base_part": p_data["base_part"],
                "rm_count": len(p_data["records"]),
                "sample_rm": p_data["records"][0]["rm_erp"]
            })
        if len(samples) >= 15:
            break
            
    stats = {
        "total_records": len(records),
        "total_unique_parts": len(parts_dict),
        "total_base_parts": len(base_parts_dict),
        "total_images_extracted": len(img_filenames),
        "records_with_images": sum(1 for r in records if r["image_url"]),
        "records_with_rm_erp": sum(1 for r in records if r["rm_erp"])
    }
    
    print("Summary Stats:", stats)
    
    output_data = {
        "stats": stats,
        "parts": parts_dict,
        "base_parts": base_parts_dict,
        "samples": samples,
        "records": records
    }
    
    with open("data_store.json", "w", encoding="utf-8") as f:
        json.dump(output_data, f, indent=2)
        
    print("Successfully updated data_store.json with complete PRINT fields!")

if __name__ == "__main__":
    extract_all()
